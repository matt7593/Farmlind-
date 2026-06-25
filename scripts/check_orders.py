import json
import base64
import urllib.request
import urllib.parse
import io
import os
from datetime import datetime

import anthropic

CLIENT_ID = os.environ["GMAIL_CLIENT_ID"]
CLIENT_SECRET = os.environ["GMAIL_CLIENT_SECRET"]
ANTHROPIC_API_KEY = os.environ["ANTHROPIC_API_KEY"]

ACCOUNTS = [
    {
        "email": "matt@farmlindproduce.com",
        "refresh_token": os.environ["GMAIL_REFRESH_TOKEN_1"]
    },
    {
        "email": "farmlindproduce@gmail.com",
        "refresh_token": os.environ["GMAIL_REFRESH_TOKEN_2"]
    }
]

VENDORS = {
    "Goodness Gardens": ["goodness"],
    "Dagele Brothers":  ["dagele"],
    "Dottavio":         ["dottavio"],
}

NOTIFY_EMAILS = ["matt@farmlindproduce.com", "farmlindproduce@gmail.com"]

# An order plus its follow-up parts are grouped into one "session" if sent within
# this many hours of each other. Avoids splitting an evening's order across the
# UTC midnight boundary (which calendar-date logic does).
SESSION_HOURS = 8

# Wait this many minutes after Matt's most recent order before sending the reminder.
# 30 min gives him time to email all three vendors before we report anything.
REMIND_DELAY_MIN = 30

REMINDER_SUBJECT = "Order Check -"  # single combined email: missing items + skipped vendors

claude = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)


def env_is_true(name):
    return os.environ.get(name, "").lower() == "true"


def get_access_token(refresh_token):
    data = urllib.parse.urlencode({
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
        "refresh_token": refresh_token,
        "grant_type": "refresh_token"
    }).encode()
    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data)
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())["access_token"]


def gmail_get(access_token, path, params=None):
    url = f"https://gmail.googleapis.com/gmail/v1{path}"
    if params:
        url += "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {access_token}"})
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())


def gmail_post(access_token, path, body):
    url = f"https://gmail.googleapis.com/gmail/v1{path}"
    data = json.dumps(body).encode()
    req = urllib.request.Request(
        url, data=data,
        headers={"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())


def decode_b64(data):
    return base64.urlsafe_b64decode(data + "==")


def extract_text_from_part(part):
    mime = part.get("mimeType", "")
    body_data = part.get("body", {}).get("data", "")
    if mime == "text/plain" and body_data:
        return decode_b64(body_data).decode("utf-8", errors="ignore")
    for sub in part.get("parts", []):
        result = extract_text_from_part(sub)
        if result:
            return result
    return ""


def extract_attachments(payload):
    attachments = []
    def walk(part):
        filename = part.get("filename", "")
        mime = part.get("mimeType", "")
        body = part.get("body", {})
        data = body.get("data", "")
        if filename and data:
            attachments.append((filename, mime, decode_b64(data), None))
        elif filename and body.get("attachmentId"):
            attachments.append((filename, mime, None, body["attachmentId"]))
        for sub in part.get("parts", []):
            walk(sub)
    walk(payload)
    return attachments


def fetch_attachment_data(access_token, message_id, attachment_id):
    result = gmail_get(access_token, f"/users/me/messages/{message_id}/attachments/{attachment_id}")
    return decode_b64(result["data"])


def parse_items_with_claude(content_blocks):
    message = claude.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=1024,
        messages=[{
            "role": "user",
            "content": content_blocks + [{
                "type": "text",
                "text": (
                    "This is a produce order. Extract every ordered item as a simple list, "
                    "one item per line, including quantities if present. "
                    "Return only the list, no other text. If there are no order items, reply NONE."
                )
            }]
        }]
    )
    text = message.content[0].text.strip()
    if text.upper() == "NONE":
        return []
    return [l.strip() for l in text.splitlines() if l.strip()]


def body_has_any_vendor_section(body_text):
    all_keywords = [kw for kws in VENDORS.values() for kw in kws]
    for line in body_text.splitlines():
        if any(kw in line.strip().lower() for kw in all_keywords):
            return True
    return False


def extract_vendor_section(body_text, vendor_keywords):
    """Extract only the lines belonging to this vendor's section of the email body."""
    lines = body_text.splitlines()
    section_lines = []
    in_section = False

    all_keywords = [kw for kws in VENDORS.values() for kw in kws]

    for line in lines:
        line_lower = line.strip().lower()

        is_our_vendor = any(kw in line_lower for kw in vendor_keywords)
        is_other_vendor = any(kw in line_lower for kw in all_keywords) and not is_our_vendor

        if is_our_vendor:
            in_section = True
            continue
        elif is_other_vendor and in_section:
            break
        elif in_section and line.strip():
            section_lines.append(line.strip())

    return "\n".join(section_lines)


def get_items_from_message(access_token, message_id, payload, vendor_keywords):
    """Get order items for a specific vendor from an email."""
    body_text = extract_text_from_part(payload)

    if body_text.strip():
        section = extract_vendor_section(body_text, vendor_keywords)
        if section.strip():
            try:
                return parse_items_with_claude([{"type": "text", "text": section}])
            except Exception as e:
                print(f"    Claude parse error: {e}")
        # Body has vendor-labeled sections but not ours — wrong email for this vendor
        if body_has_any_vendor_section(body_text):
            return []

    # Body is empty or has no vendor labels — try attachments
    items = []
    for att in extract_attachments(payload):
        filename, mime, data, att_id = att
        if att_id:
            data = fetch_attachment_data(access_token, message_id, att_id)
        print(f"    Processing attachment: {filename}")
        try:
            content_block = None
            if mime == "application/pdf" or filename.lower().endswith(".pdf"):
                content_block = {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()}}
            elif mime.startswith("image/"):
                content_block = {"type": "image", "source": {"type": "base64", "media_type": mime, "data": base64.b64encode(data).decode()}}
            elif filename.lower().endswith((".xlsx", ".xls")):
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
                text = "\n".join(" ".join(str(c) for c in row if c is not None) for sheet in wb.worksheets for row in sheet.iter_rows(values_only=True))
                content_block = {"type": "text", "text": text}
            elif filename.lower().endswith((".docx", ".doc")):
                from docx import Document
                doc = Document(io.BytesIO(data))
                content_block = {"type": "text", "text": "\n".join(p.text for p in doc.paragraphs if p.text.strip())}
            else:
                try:
                    content_block = {"type": "text", "text": data.decode("utf-8", errors="ignore")}
                except Exception:
                    pass
            if content_block:
                items.extend(parse_items_with_claude([content_block]))
        except Exception as e:
            print(f"    Attachment error: {e}")

    return items


ORDER_RECIPIENTS = "to:orders@goodnessgardens.com OR to:Office@dagelebrothersproduce.com OR to:anthony@mdottavioproduce.com"

VENDOR_EMAILS = {
    "Goodness Gardens": "orders@goodnessgardens.com",
    "Dagele Brothers":  "office@dagelebrothersproduce.com",
    "Dottavio":         "anthony@mdottavioproduce.com",
}


def msg_sent_to_vendor(msg, vendor_name):
    """Return True if the message's To header matches this vendor's email address."""
    vendor_addr = VENDOR_EMAILS.get(vendor_name, "").lower()
    if not vendor_addr:
        return False
    headers = msg.get("payload", {}).get("headers", [])
    for h in headers:
        if h.get("name", "").lower() == "to":
            if vendor_addr in h.get("value", "").lower():
                return True
    return False


def fetch_all_recent_orders(access_tokens, days=16, max_results=60):
    """Return [(ts_ms, access_token, full_message)] for every order email sent in the
    last `days`, across all accounts, sorted oldest-first. Uses internalDate (epoch ms,
    timezone-independent) so there are no calendar-boundary issues."""
    out = []
    for email_addr, token in access_tokens:
        try:
            result = gmail_get(token, "/users/me/messages", {
                "q": f"in:sent ({ORDER_RECIPIENTS}) newer_than:{days}d",
                "maxResults": max_results
            })
            for m in result.get("messages", []):
                full = gmail_get(token, f"/users/me/messages/{m['id']}", {"format": "full"})
                out.append((int(full.get("internalDate", 0)), token, full))
        except Exception as e:
            print(f"  Error fetching orders for {email_addr}: {e}")
    out.sort(key=lambda x: x[0])
    return out


def reminder_already_sent_after(access_token, since_ms, subject_prefix):
    """Return True if an email with the given subject prefix was already sent after since_ms."""
    since_dt = datetime.fromtimestamp(since_ms / 1000)
    after = since_dt.strftime("%Y/%m/%d")
    try:
        result = gmail_get(access_token, "/users/me/messages", {
            "q": f"in:sent subject:\"{subject_prefix}\" after:{after}",
            "maxResults": 10
        })
        for m in result.get("messages", []):
            full = gmail_get(access_token, f"/users/me/messages/{m['id']}", {"format": "metadata"})
            if int(full.get("internalDate", 0)) > since_ms:
                return True
    except Exception:
        pass
    return False


def collect_vendor_items(messages, keywords):
    """Given a list of (ts, token, full_message), return all order items for this vendor,
    deduplicating identical vendor sections that appear in more than one account."""
    items = []
    seen_sections = set()
    for ts, token, msg in messages:
        body = extract_text_from_part(msg.get("payload", {}))
        section = extract_vendor_section(body, keywords).strip()
        if section and section in seen_sections:
            continue
        msg_items = get_items_from_message(token, msg["id"], msg.get("payload", {}), keywords)
        if msg_items:
            if section:
                seen_sections.add(section)
            items.extend(msg_items)
    return items


def item_matches(prev_name, today_name):
    """True if these two normalized item names refer to the same product, tolerant of
    typos and minor spelling differences (e.g. 'oregeno' vs 'oregano')."""
    import difflib
    # Exact substring match (handles 'green bell' vs 'jumbo green bell')
    if prev_name in today_name or today_name in prev_name:
        return True
    # Whole-string fuzzy match for typos
    if difflib.SequenceMatcher(None, prev_name, today_name).ratio() >= 0.82:
        return True
    # Word-level fuzzy match: every word in the shorter name has a close match
    # in the longer one (handles a typo inside a multi-word item)
    short, long = sorted([prev_name.split(), today_name.split()], key=len)
    if short and all(
        any(difflib.SequenceMatcher(None, w, lw).ratio() >= 0.85 for lw in long)
        for w in short
    ):
        return True
    return False


def normalize_items(items):
    """Ask Claude to return just the product names, no quantities."""
    response = claude.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=1024,
        messages=[{
            "role": "user",
            "content": (
                "Strip all leading quantities and numbers from these produce order items and return only the product names, one per line. "
                "Keep descriptors like colors, sizes, and grades (e.g. 'jumbo', 'red', 'fancy', '60ct'). "
                "No bullet points, no numbers at the start, no extra text.\n\n"
                + "\n".join(items)
            )
        }]
    )
    return [l.strip().lower() for l in response.content[0].text.strip().splitlines() if l.strip()]



def build_reminder_body(missing_by_vendor, vendor_today, skipped_vendors=None):
    lines = [
        "Hi Matt,",
        "",
        "Here is a summary of tonight's orders and what may be missing:",
        "",
        "=" * 40,
    ]
    for vendor, missing in missing_by_vendor.items():
        sent = vendor_today.get(vendor, [])
        lines.append("")
        lines.append(f"  {vendor.upper()}")
        lines.append(f"  {'-' * 36}")
        lines.append(f"  This week you sent these items:")
        for item in sent:
            lines.append(f"    - {item}")
        lines.append("")
        if missing:
            lines.append(f"  You may have forgotten these items ({len(missing)}):")
            for item in missing:
                lines.append(f"    - {item}")
        else:
            lines.append(f"  No missing items.")
        lines.append("")
        lines.append("=" * 40)
    if skipped_vendors:
        lines.append("")
        lines.append("  *** YOU DID NOT SEND AN ORDER TO: ***")
        for vendor in skipped_vendors:
            lines.append(f"    *** {vendor.upper()} ***")
        lines.append("")
        lines.append("  If you skipped them on purpose, ignore this.")
        lines.append("  Otherwise you may have forgotten to email them.")
        lines.append("")
        lines.append("=" * 40)
    lines += ["", "- Farmlind Order Bot"]
    return "\n".join(lines)


def preview_reminder(missing_by_vendor, vendor_today, skipped_vendors, today):
    print(f"Subject: {REMINDER_SUBJECT} {today.strftime('%A %b %d')}")
    print(f"To: {', '.join(NOTIFY_EMAILS)}")
    print(build_reminder_body(missing_by_vendor, vendor_today, skipped_vendors))


def send_reminder(access_token, sender_email, missing_by_vendor, vendor_today, skipped_vendors, today):
    subject = f"{REMINDER_SUBJECT} {today.strftime('%A %b %d')}"
    body = build_reminder_body(missing_by_vendor, vendor_today, skipped_vendors)
    raw = f"From: {sender_email}\r\nTo: {', '.join(NOTIFY_EMAILS)}\r\nSubject: {subject}\r\n\r\n{body}"
    encoded = base64.urlsafe_b64encode(raw.encode()).decode()
    gmail_post(access_token, "/users/me/messages/send", {"raw": encoded})
    print("Reminder sent.")


STATE_FILE = ".order_state.json"


def load_state():
    try:
        with open(STATE_FILE) as f:
            return json.load(f)
    except Exception:
        return {}


def save_state(state):
    try:
        with open(STATE_FILE, "w") as f:
            json.dump(state, f)
    except Exception:
        pass


def main(override_date=None):
    today = override_date if override_date else datetime.now()
    print(f"Checking orders for {today.strftime('%Y-%m-%d %A')}")

    access_tokens = []
    for account in ACCOUNTS:
        try:
            token = get_access_token(account["refresh_token"])
            access_tokens.append((account["email"], token))
            print(f"Authenticated: {account['email']}")
        except Exception as e:
            print(f"Auth failed for {account['email']}: {e}")

    if not access_tokens:
        print("No valid tokens, exiting.")
        return

    session_ms = SESSION_HOURS * 3600 * 1000
    now_ms = int(datetime.now().timestamp() * 1000)

    all_orders = fetch_all_recent_orders(access_tokens)
    if not all_orders:
        print("No recent order emails found — nothing to do.")
        return

    # Anchor the current order session on the most recent order.
    if override_date:
        cutoff = int(datetime.combine(override_date.date(), datetime.max.time()).timestamp() * 1000)
        candidates = [o for o in all_orders if o[0] <= cutoff]
        if not candidates:
            print("No order on or before that date — nothing to do.")
            return
        anchor_ts = candidates[-1][0]
    else:
        anchor_ts = all_orders[-1][0]

    anchor_age_min = (now_ms - anchor_ts) / 60000
    sender_email, sender_token = access_tokens[0]
    force = bool(override_date) or env_is_true("FORCE_SEND")
    dry = env_is_true("DRY_RUN")

    # Wait 30 min after the most recent order so Matt has time to email all vendors
    # before we report anything (he often sends them a few minutes apart).
    if not force and anchor_age_min < REMIND_DELAY_MIN:
        print(f"Most recent order sent {anchor_age_min:.1f} min ago — waiting until {REMIND_DELAY_MIN} min mark.")
        return

    session_start = anchor_ts - session_ms
    session_end = anchor_ts + session_ms
    current_msgs = [o for o in all_orders if session_start <= o[0] <= session_end]
    prior_msgs = [o for o in all_orders if o[0] < session_start]
    print(f"Session anchored at {datetime.fromtimestamp(anchor_ts/1000)}; "
          f"{len(current_msgs)} order email(s) in session, {len(prior_msgs)} prior. "
          f"Order is {anchor_age_min:.0f} min old.")

    # --- Cheap detection (no Claude): which vendors were ordered, which were skipped ---
    vendor_has_today = {}
    skipped_vendors = []
    for vendor_name, keywords in VENDORS.items():
        has_today = any(
            extract_vendor_section(extract_text_from_part(msg.get("payload", {})), keywords).strip()
            or msg_sent_to_vendor(msg, vendor_name)
            for ts, token, msg in current_msgs
        )
        vendor_has_today[vendor_name] = has_today
        if not has_today:
            ordered_before = any(
                extract_vendor_section(extract_text_from_part(msg.get("payload", {})), keywords).strip()
                or msg_sent_to_vendor(msg, vendor_name)
                for ts, token, msg in prior_msgs
            )
            if ordered_before:
                skipped_vendors.append(vendor_name)

    if not any(vendor_has_today.values()):
        print("\nNo orders in this session — nothing to compare.")
        return

    # --- Check if we already sent the reminder for this anchor ---
    state = load_state()
    key = str(anchor_ts)
    st = state.get(key, {})

    already_sent = bool(st.get("sent"))
    if not force and not already_sent:
        already_sent = reminder_already_sent_after(sender_token, anchor_ts, REMINDER_SUBJECT)

    if already_sent and not force:
        print("Reminder already sent for this order session — nothing to do.")
        return

    # --- Run item comparison via Claude ---
    vendor_today = {}
    missing_by_vendor = {}
    for vendor_name, keywords in VENDORS.items():
        if not vendor_has_today[vendor_name]:
            vendor_today[vendor_name] = []
            continue
        print(f"\nChecking {vendor_name}...")
        today_items = collect_vendor_items(current_msgs, keywords)
        vendor_today[vendor_name] = today_items

        prev_items = []
        prev_anchor_ts = None
        for ts, token, msg in reversed(prior_msgs):
            if (extract_vendor_section(extract_text_from_part(msg.get("payload", {})), keywords).strip()
                    or msg_sent_to_vendor(msg, vendor_name)):
                prev_anchor_ts = ts
                break
        if prev_anchor_ts is not None:
            p_start = prev_anchor_ts - session_ms
            p_end = prev_anchor_ts + session_ms
            prev_session = [o for o in prior_msgs if p_start <= o[0] <= p_end]
            prev_items = collect_vendor_items(prev_session, keywords)
        print(f"  Today: {len(today_items)} items, Previous: {len(prev_items)} items")

        if today_items and prev_items:
            try:
                today_names = normalize_items(today_items)
                prev_names = normalize_items(prev_items)
                missing, seen = [], set()
                for name in prev_names:
                    if name in seen:
                        continue
                    seen.add(name)
                    if not any(item_matches(name, tn) for tn in today_names):
                        missing.append(name)
                if missing:
                    missing_by_vendor[vendor_name] = missing
            except Exception as e:
                print(f"  Comparison error for {vendor_name}: {e}")

    print(f"\nMissing items: {missing_by_vendor}")
    print(f"Skipped vendors: {skipped_vendors}")

    # Send one combined email covering missing items + any skipped vendors
    has_something_to_report = bool(missing_by_vendor) or bool(skipped_vendors)
    if has_something_to_report:
        if dry:
            print("\n--- DRY RUN: email NOT sent. Preview ---")
            preview_reminder(missing_by_vendor, vendor_today, skipped_vendors, today)
        else:
            send_reminder(sender_token, sender_email, missing_by_vendor, vendor_today, skipped_vendors, today)
    else:
        print("No missing items and no skipped vendors — nothing to report.")

    if not dry:
        st["sent"] = True
        state[key] = st
        save_state(state)


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        override = datetime.strptime(sys.argv[1], "%Y-%m-%d")
        main(override_date=override)
    else:
        main()
